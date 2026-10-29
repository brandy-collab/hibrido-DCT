#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
HEV P2 + 6-speed DCT Simulator (Kia Niro HEV / Hyundai "TMED" layout)
=====================================================================

Parallel P2 hybrid modelled after MathWorks' "HEV P2 Reference Application"
(engine -> disconnect clutch -> motor -> dual-clutch transmission -> wheels,
ECMS supervisory control, speed/load shift schedule), parameterised as a
2nd-generation Kia Niro HEV:

  engine   1.6 GDi Atkinson, 77.2 kW @ 5700 rpm, 144 Nm @ 4000 rpm
  HSG      belt starter-generator, 8 kW / 35.3 Nm (engine start, speed sync,
           standstill charging)
  K0       engine clutch between engine and motor
  motor    P2 PMSM on the DCT input, 32 kW / 170 Nm
  6DCT     dry dual-clutch: 3.600 2.217 1.371 0.953 0.978 0.767,
           final drives 4.438 (gears 1-4) / 3.227 (gears 5-6)
  battery  Li-ion polymer, 240 V, 1.32 kWh, 42 kW

Model
-----
  Plant: up to three rotating bodies -- engine (+HSG), DCT input (motor +
  input shafts), vehicle (wheel) -- coupled through K0 and the active DCT
  clutch, each either locked or slipping. Gear shifts are clutch-to-clutch
  with an inertia phase: the clutch whose slip matches the sign of the driver
  demand carries the torque while the motor synchronises the input speed;
  the on-coming clutch locks at synchronism. Slip losses are accounted for.
  VCU: ECMS -- each 0.1 s it compares EV against engine-coupled torque splits
  on  J = m_fuel + s * P_batt,chem / LHV,  with an SOC-feedback equivalence
  factor s, plus on/off hysteresis and minimum on/off times. Series
  regenerative braking. Engine start: HSG cranks, ECU syncs engine to the
  motor speed, K0 closes. Standstill charging via HSG below the coupling speed.
  TCU: up/down-shift speeds interpolated between light- and full-pedal tables,
  engine-protect downshifts, minimum time between shifts.

Shared component models (mapped engine, loss-map machines, Li-ion pack,
drive cycles, gauges) come from hev_ips_sim.py, the power-split version.

Usage
-----
  python hev_p2dct_sim.py --cycle WLTC
  python hev_p2dct_sim.py --cycle UDDS --set mass=1550 --export out/
  python hev_p2dct_sim.py --list-params
"""

import argparse
import csv
import math
import os
from dataclasses import dataclass, fields, asdict

import numpy as np

import hev_ips_sim as common
from hev_ips_sim import (RPM, G, interp1, Machine, Battery, make_scenario, CYCLES,
                         C_ENG, C_MG1, C_MG2, C_BAT, C_BRK, C_REF, C_TXT, C_TXT2, C_GRID, C_SURF)

C_HSG = C_MG1                       # HSG uses the violet slot
MODE_NAMES = ["Stationary", "EV", "HEV · motor charging", "HEV · motor assist", "Idle charging"]
MODE_COLORS = ["#c9c8c2", "#1baf7a", "#2a78d6", "#eb6834", "#6250d6"]
ENG_STATES = ["OFF", "START", "RUN", "CHARGE", "STOP"]


# =============================================================================
# Parameters (Kia Niro HEV, 2nd generation)
# =============================================================================
@dataclass
class P2Params:
    # --- Vehicle body ---
    mass: float = 1500.0          # kg, kerb + driver
    Cd: float = 0.29
    A: float = 2.36               # m2
    Crr: float = 0.0085
    r_wheel: float = 0.320        # m, 205/60 R16 rolling radius
    J_wheels: float = 3.2
    # --- 6-speed DCT ---
    g1: float = 3.600
    g2: float = 2.217
    g3: float = 1.371
    g4: float = 0.953
    g5: float = 0.978
    g6: float = 0.767
    fd_odd_low: float = 4.438     # final drive, gears 1-4
    fd_high: float = 3.227        # final drive, gears 5-6
    gb_eta: float = 0.965         # load efficiency (mesh + bearings)
    gb_spin_Nm: float = 0.8       # drag at input shaft
    gb_visc: float = 0.002        # Nm/(rad/s) at input shaft
    input_J: float = 0.04         # DCT input shafts + clutch discs
    dct_Tin_max: float = 265.0    # Nm, DCT input torque limit (= Kia system torque)
    shift_min_s: float = 0.8      # min. time between shifts
    shift_tau: float = 0.10       # s, input-speed synchronisation time constant
    # shift schedule: upshift speed n -> n+1 at light / full pedal [km/h]
    up12_lo: float = 14.0
    up23_lo: float = 24.0
    up34_lo: float = 36.0
    up45_lo: float = 50.0
    up56_lo: float = 62.0
    up12_hi: float = 40.0
    up23_hi: float = 65.0
    up34_hi: float = 95.0
    up45_hi: float = 120.0
    up56_hi: float = 140.0
    dn_hyst_kmh: float = 7.0
    # --- Engine (Smartstream G1.6 GDi) ---
    eng_disp_L: float = 1.58
    eng_Pmax_kW: float = 77.2
    eng_rpm_idle: float = 1000.0
    eng_rpm_max: float = 6000.0
    eng_J: float = 0.16
    eng_eta_ind: float = 0.455
    eng_tau: float = 0.12
    eng_couple_rpm: float = 1100.0   # min. input speed for engine coupling
    # --- HSG (hybrid starter-generator) ---
    hsg_Pmax_kW: float = 8.0
    hsg_Tmax: float = 35.3
    hsg_ratio: float = 2.5
    hsg_rpm_max: float = 16000.0
    hsg_J: float = 0.008
    # --- P2 traction motor ---
    mot_Pmax_kW: float = 32.0
    mot_Tmax: float = 170.0
    mot_rpm_max: float = 6600.0
    mot_J: float = 0.06
    # --- Battery ---
    batt_cells: float = 64
    batt_Ah: float = 5.5
    batt_Rcell_mOhm: float = 1.6
    batt_Pdis_kW: float = 42.0
    batt_Pchg_kW: float = 36.0
    soc_init: float = 0.60
    dcdc_eff: float = 1.0         # 1 = motor inverter on the battery bus
    aux_kW: float = 0.40
    # --- Supervisory control (ECMS) ---
    soc_target: float = 0.60
    soc_low: float = 0.40         # force engine / idle charging below
    soc_ke_drop: float = 0.12     # SOC target reduction at 130 km/h (room for regen)
    ecms_s0: float = 2.6          # nominal equivalence factor
    ecms_Kp: float = 20.0         # per unit SOC error
    ecms_Ki: float = 0.02         # per unit SOC error per s
    ecms_hyst_gps: float = 0.08   # on/off cost margin [g/s]
    ecms_Te_min_frac: float = 0.15
    ev_vmax_kmh: float = 110.0
    eng_min_on_s: float = 3.0
    eng_min_off_s: float = 1.5
    eng_P_rate: float = 400.0     # Nm/s engine torque ramp (blending)
    idle_chg_kW: float = 5.0
    idle_chg_rpm: float = 1400.0
    regen_fade_kmh: float = 6.0
    # --- Driver ---
    drv_Kp: float = 1.0
    drv_Ki: float = 0.15
    drv_preview: float = 0.4
    # --- Environment ---
    grade_pct: float = 0.0
    wind_kmh: float = 0.0
    T_amb_C: float = 25.0
    p_amb_kPa: float = 101.325
    # --- Fuel ---
    fuel_LHV_MJkg: float = 43.0
    fuel_rho: float = 0.745
    co2_per_kg: float = 3.09
    # --- Scenario helpers / solver ---
    const_kmh: float = 90.0
    const_hold_s: float = 300.0
    dt: float = 0.02
    log_dt: float = 0.1


PARAM_HELP = {
    "eng_couple_rpm": "lowest input-shaft speed at which K0 may close",
    "ecms_s0": "equivalence factor converting battery energy to fuel",
    "ecms_Kp": "SOC feedback on the equivalence factor",
    "ecms_hyst_gps": "engine on/off decision margin",
    "shift_tau": "time constant of the motor's input-speed synchronisation",
    "dcdc_eff": "1.0 = no boost converter",
    "dct_Tin_max": "combined engine + motor torque the DCT accepts",
    "soc_ke_drop": "lower SOC target at speed, scaled with kinetic energy",
}


class KiaEngine(common.Engine):
    """Smartstream G1.6 GDi full-load torque (77.2 kW @ 5700, 144 Nm @ 4000)."""
    _RPM = [1000, 1500, 2000, 2500, 3000, 3500, 4000, 4500, 5000, 5500, 5700, 6000]
    _TQ = [105, 120, 130, 136, 140, 142, 144, 143, 139, 133, 129, 118]
    _P_REF = 77.2
    _RPM_REF = 6000.0


def gear_ratios(p):
    g = [p.g1, p.g2, p.g3, p.g4, p.g5, p.g6]
    fd = [p.fd_odd_low] * 4 + [p.fd_high] * 2
    return [gi * fi for gi, fi in zip(g, fd)]


def shift_speeds(p, pedal):
    """Upshift speeds [km/h] n->n+1 for n = 1..5 at pedal position."""
    lo = [p.up12_lo, p.up23_lo, p.up34_lo, p.up45_lo, p.up56_lo]
    hi = [p.up12_hi, p.up23_hi, p.up34_hi, p.up45_hi, p.up56_hi]
    f = min(max((pedal - 0.2) / 0.6, 0.0), 1.0)
    return [a + (b - a) * f for a, b in zip(lo, hi)]


SIGNALS = ["t", "v_ref", "v", "accel", "brake", "mode", "gear", "shifting", "k0", "eng_state",
           "w_e", "w_i", "Te", "Tm", "T_hsg", "T_fb", "Tc", "P_eng", "P_mot", "P_hsg",
           "P_mot_e", "P_hsg_e", "P_batt", "P_wheel", "soc", "I_batt", "V_batt", "mdot",
           "fuel_L", "dist", "P_dis_lim", "P_chg_lim", "ecms_s", "eta_eng", "grade"]


class Results:
    pass


# =============================================================================
# Simulation
# =============================================================================
def simulate(p: P2Params, scen, progress=None):
    eng = KiaEngine(p)
    mot = Machine(p.mot_Pmax_kW, p.mot_Tmax, p.mot_rpm_max, p.mot_J)
    hsg = Machine(p.hsg_Pmax_kW, p.hsg_Tmax, p.hsg_rpm_max, p.hsg_J)
    bat = Battery(p)

    R_g = gear_ratios(p)
    rw = p.r_wheel
    Gh = p.hsg_ratio
    JE = p.eng_J + hsg.J * Gh * Gh
    JI = mot.J + p.input_J
    JV = p.mass * rw * rw + p.J_wheels
    eta = p.gb_eta
    rho = p.p_amb_kPa * 1e3 / (287.05 * (p.T_amb_C + 273.15))
    cdA = 0.5 * rho * p.Cd * p.A
    v_wind = p.wind_kmh / 3.6
    eta_dc = p.dcdc_eff
    P_aux = p.aux_kW * 1e3
    LHV = eng.LHV
    w_couple = p.eng_couple_rpm * RPM
    w_chg = p.idle_chg_rpm * RPM
    w_emax = eng.w_max

    dt = p.dt
    n_steps = int(scen.duration / dt) + 1
    log_every = max(1, int(round(p.log_dt / dt)))
    ecms_every = max(1, int(round(0.1 / dt)))
    t_cyc, v_cyc, g_cyc = list(scen.t), list(scen.v), list(scen.grade_pct)
    T_brake_max = 1.0 * p.mass * G * rw
    M_eq_drv = (JV + JI * R_g[0] ** 2 * 0.3) / rw ** 2     # rough, driver feed-forward

    # ---------------- states
    we = wi = ww = 0.0
    Te = 0.0
    soc = p.soc_init
    drv_I = 0.0
    s_I = 0.0
    gear = 1
    shift_to = 0                  # 0 = no shift in progress
    shift_t = 0.0
    last_shift = -10.0
    n_shifts = 0
    k0 = False                    # engine clutch locked?
    eng_state = "OFF"
    eng_timer = 0.0               # time in current on/off condition
    ecms_on = False
    ecms_Te = 0.0
    ecms_s = p.ecms_s0
    want_couple = False
    dist = 0.0
    fuel_kg = 0.0
    n_starts = 0
    wot_done_t = None
    t100 = None
    a_w_prev = 0.0
    # cycle lookups precomputed once (driver reference, preview, feed-forward accel, grade)
    t_steps = np.arange(n_steps) * dt
    VR = np.interp(t_steps, scen.t, scen.v)
    VT = np.interp(t_steps + p.drv_preview, scen.t, scen.v)
    AR = (np.interp(t_steps + p.drv_preview + 0.5, scen.t, scen.v)
          - np.interp(t_steps + p.drv_preview - 0.5, scen.t, scen.v))
    GR = np.interp(t_steps, scen.t, scen.grade_pct)
    VR, VT, AR, GR = VR.tolist(), VT.tolist(), AR.tolist(), GR.tolist()
    P_dis_lim, P_chg_lim = bat.limits(soc)

    E = {k: 0.0 for k in ["fuel", "eng_mech", "eng_loss", "mot_mech", "mot_elec", "mot_loss",
                          "hsg_mech", "hsg_elec", "hsg_loss", "dcdc_loss", "batt_chem", "batt_loss",
                          "batt_out", "batt_in", "aux", "aero", "roll", "grade", "brake",
                          "gearbox", "slip_dct", "lock", "standstill",
                          "mot_motor_mech", "mot_motor_elec", "mot_gen_mech", "mot_gen_elec"]}
    mode_time = [0.0] * 5
    gear_time = [0.0] * 6
    eng_eff = []
    log = {s: [] for s in SIGNALS}

    def ke():
        return 0.5 * (JE * we * we + JI * wi * wi + JV * ww * ww)

    def ecms(wi_c, T_tot, soc_c, s_c, P_dis, P_chg):
        """Return (J_ev or inf, J_hev_best or inf, Te_best)."""
        V, Rb = bat.ocv(soc_c), bat.r(soc_c)
        Tm_lim = mot.tlim(wi_c)

        def cost_elec(Tm):
            Pe = Tm * wi_c + (mot.c0 + mot.kc * Tm * Tm + mot.ki * wi_c + mot.kw * wi_c ** 3)
            Pb = Pe + P_aux
            Pb = np.where(Pb > 0, Pb / eta_dc, Pb * eta_dc)
            ok = (Pb <= P_dis) & (Pb >= -P_chg)
            disc = np.maximum(V * V - 4 * Rb * Pb, 0.0)
            I = (V - np.sqrt(disc)) / (2 * Rb)
            return s_c * V * I / LHV, ok
        # EV
        Tm_ev = np.array([T_tot])
        c_ev, ok_ev = cost_elec(Tm_ev)
        J_ev = float(c_ev[0]) if (ok_ev[0] and abs(T_tot) <= Tm_lim) else math.inf
        # HEV
        J_h, Te_b = math.inf, 0.0
        if w_couple <= wi_c <= w_emax:
            tmx = eng.tmax(wi_c)
            Teg = np.linspace(p.ecms_Te_min_frac * tmx, tmx, 24)
            Tm = T_tot - Teg
            c_el, ok = cost_elec(Tm)
            ok &= np.abs(Tm) <= Tm_lim
            mf = eng.fuel_rate_v(np.full_like(Teg, wi_c), Teg)
            J = mf + c_el
            if np.any(ok):
                i = int(np.argmin(np.where(ok, J, np.inf)))
                J_h, Te_b = float(J[i]), float(Teg[i])
            else:      # demand beyond capability: engine flat out, motor at its limit
                J_h, Te_b = 1.0, float(tmx)
        return J_ev, J_h, Te_b

    for step in range(n_steps):
        t = step * dt
        v = ww * rw
        grade = GR[step]
        th = math.atan(grade / 100.0)
        if step % ecms_every == 0:
            P_dis_lim, P_chg_lim = bat.limits(soc)
        R_now = R_g[gear - 1]
        R_tgt = R_g[(shift_to or gear) - 1]
        wi_tgt = R_tgt * ww                       # synchronous input speed

        # ================================================= DRIVER
        Tw_max = min(eng.tmax(max(wi_tgt, w_couple)) + mot.tlim(max(wi_tgt, 1.0)),
                     p.dct_Tin_max) * R_tgt * eta
        if scen.kind == "wot":
            if wot_done_t is None and v >= 130 / 3.6:
                wot_done_t = t
            if wot_done_t is None and t < 35:
                accel, brake = 1.0, 0.0
            else:
                accel, brake = 0.0, 0.25
            v_r = v
        else:
            v_r, v_tg, a_r = VR[step], VT[step], AR[step]
            F_ff = (p.mass * 1.04 * a_r + cdA * (v_tg + v_wind) * abs(v_tg + v_wind)
                    + (p.Crr * p.mass * G * math.cos(th) if v_tg > 0.05 else 0.0)
                    + p.mass * G * math.sin(th))
            err = v_tg - v
            Tw_cmd = (F_ff + p.drv_Kp * p.mass * err + drv_I) * rw
            hold = v_tg < 0.05 and v < 0.3
            if hold:
                accel, brake = 0.0, 0.3
                drv_I = 0.0
            elif Tw_cmd >= 0:
                accel, brake = min(Tw_cmd / Tw_max, 1.0), 0.0
            else:
                accel, brake = 0.0, min(-Tw_cmd / T_brake_max, 1.0)
            sat = (accel >= 1.0 and err > 0) or (brake >= 1.0 and err < 0)
            if not sat and not hold:
                drv_I += p.drv_Ki * p.mass * err * dt
        if t100 is None and v >= 100 / 3.6:
            t100 = t
        Tw_req = accel * Tw_max - brake * T_brake_max

        # ================================================= TCU (shift schedule)
        if not shift_to and t - last_shift >= p.shift_min_s:
            ups = shift_speeds(p, 0.0 if brake > 0 else accel)
            vk = v * 3.6
            new = gear
            if gear < 6 and vk > ups[gear - 1]:
                new = gear + 1
                if k0 and R_g[new - 1] * ww < w_couple * 1.05:
                    new = gear                     # keep engine above coupling speed
            elif gear > 1 and vk < ups[gear - 2] - p.dn_hyst_kmh:
                new = gear - 1
            if k0 and gear > 1 and R_now * ww < w_couple * 1.05:
                new = gear - 1                     # engine-protect downshift
            if R_g[gear - 1] * ww > w_emax * 0.96 and gear < 6:
                new = gear + 1                     # over-speed upshift
            if new != gear:
                shift_to, shift_t = new, 0.0
                n_shifts += 1
                R_tgt = R_g[new - 1]
                wi_tgt = R_tgt * ww

        # input-side torque the driver asks for (through the target ratio)
        spin = p.gb_spin_Nm + p.gb_visc * wi
        if Tw_req >= 0:
            T_tot = min((Tw_req / R_tgt + spin) / eta, p.dct_Tin_max)
        else:
            T_tot = (Tw_req / R_tgt + spin) / (2 - eta)

        # ================================================= VCU: ECMS (10 Hz)
        if step % ecms_every == 0:
            soc_tgt = p.soc_target - p.soc_ke_drop * (v / (130 / 3.6)) ** 2
            e_soc = soc_tgt - soc
            s_I = min(max(s_I + p.ecms_Ki * e_soc * 0.1, -1.0), 1.0)
            ecms_s = min(max(p.ecms_s0 + p.ecms_Kp * e_soc + s_I, 1.0), 6.0)
            couple_ok = w_couple <= wi_tgt <= w_emax
            if Tw_req > 0 and couple_ok:
                J_ev, J_h, Te_b = ecms(wi_tgt, T_tot, soc, ecms_s, P_dis_lim, P_chg_lim)
                hy = p.ecms_hyst_gps / 1000.0
                forced = soc < p.soc_low or v * 3.6 > p.ev_vmax_kmh or J_ev == math.inf
                if eng_state in ("RUN", "START") and want_couple:
                    on = forced or J_h < J_ev + hy or eng_timer < p.eng_min_on_s
                else:
                    on = forced or (J_h + hy < J_ev and eng_timer >= p.eng_min_off_s)
                ecms_on, ecms_Te = on, Te_b
            else:
                ecms_on, ecms_Te = False, 0.0
                if Tw_req > 0 and eng_state == "RUN" and k0 and eng_timer < p.eng_min_on_s \
                        and couple_ok:
                    ecms_on = True                 # respect min on-time while driving
                    ecms_Te = max(min(T_tot, eng.tmax(wi_tgt)), p.ecms_Te_min_frac * eng.tmax(wi_tgt))

        idle_chg = (soc < p.soc_low + 0.02 and wi_tgt < w_couple and not (Tw_req < 0 and v > 3)) \
            or (eng_state == "CHARGE" and soc < p.soc_low + 0.05 and wi_tgt < w_couple)
        want_couple = ecms_on
        want_engine = want_couple or idle_chg

        # ================================================= engine state machine
        prev_state = eng_state
        if eng_state in ("OFF", "STOP"):
            if want_engine:
                eng_state = "START"; n_starts += 1
        elif eng_state == "START":
            if not want_engine:
                eng_state = "STOP"
            elif want_couple and we > 0.7 * w_couple and abs(we - wi) < 3.0 and not shift_to:
                KE_pre = ke(); k0 = True; eng_state = "RUN"
                we = wi; E["lock"] += KE_pre - ke()
            elif not want_couple and idle_chg and abs(we - w_chg) < 15:
                eng_state = "CHARGE"
        elif eng_state == "RUN":
            if (not want_couple) or wi < 0.9 * w_couple or Tw_req < 0:
                k0 = False
                eng_state = "CHARGE" if (idle_chg and Tw_req >= 0) else "STOP"
        elif eng_state == "CHARGE":
            if want_couple:
                eng_state = "START"                # resync and couple
            elif not idle_chg:
                eng_state = "STOP"
        on_now = eng_state in ("START", "RUN", "CHARGE")
        if (prev_state in ("START", "RUN", "CHARGE")) != on_now:
            eng_timer = 0.0
        else:
            eng_timer += dt

        # ================================================= engine / HSG torque commands
        Tf = eng.tfric(we)
        T_hsg = 0.0
        fuel_on = False
        if eng_state == "START":
            w_sync = wi if want_couple else w_chg
            if we < 0.7 * max(min(w_sync, w_couple), eng.w_idle * 0.9):
                T_hsg = hsg.tlim(Gh * we)                     # crank
            else:
                fuel_on = True
                Te_cmd = Tf * 0 + JE * (w_sync - we) / 0.15 + JE * a_w_prev * R_tgt
                Te = min(max(Te_cmd, 0.0), eng.tmax(we))
                if Te_cmd < 0:                                # HSG helps pull speed down
                    T_hsg = max(Te_cmd / Gh, -hsg.tlim(Gh * we))
        elif eng_state == "RUN":
            fuel_on = True
            Te_tgt = min(ecms_Te, eng.tmax(we))
            Te += min(max(Te_tgt - Te, -p.eng_P_rate * dt), p.eng_P_rate * dt)
        elif eng_state == "CHARGE":
            fuel_on = True
            T_hsg = -min(p.idle_chg_kW * 1e3 / max(Gh * we, 1.0), hsg.tlim(Gh * we))
            Te_cmd = -T_hsg * Gh + JE * (w_chg - we) / 0.2
            Te = min(max(Te_cmd, 0.0), eng.tmax(we))
        if not fuel_on:
            Te = 0.0
        Te_net = Te if fuel_on else -Tf
        if eng_state == "STOP" and we < 4.0:
            eng_state = "OFF"
        Te_c = Te_net if k0 else 0.0              # engine torque reaching the input shaft

        # ================================================= motor torque & DCT clutch
        Tm_lim = mot.tlim(wi)
        dct_lock = shift_to == 0
        Tc = 0.0
        R_cl = R_now
        if dct_lock:
            if Tw_req >= 0:
                Tm = T_tot - Te_c
            else:
                fade = min(max((v * 3.6 - 1.5) / max(p.regen_fade_kmh - 1.5, 0.1), 0.0), 1.0)
                Tm = T_tot * fade - Te_c
        else:
            shift_t += dt
            # the clutch whose slip sign agrees with the demand carries the torque
            slip_new = wi - R_tgt * ww
            dem = 1.0 if T_tot > 0.5 else (-1.0 if T_tot < -0.5 else 0.0)
            if dem == 0.0 or (slip_new > 0) == (dem > 0):
                R_cl = R_tgt
            else:
                R_cl = R_now
            T_cl_dem = (Tw_req / R_cl) * (1 / eta if Tw_req >= 0 else eta)
            slip = wi - R_cl * ww
            # clutch torque during the inertia phase is limited to what the input
            # side (motor + coupled engine) can react; brakes cover the rest
            Tc_cap = 0.85 * mot.tlim(wi) + abs(Te_c)
            Tc_mag = min(abs(T_cl_dem), Tc_cap)
            Tc = Tc_mag * (1.0 if slip > 0 else -1.0) if abs(slip) > 1e-6 else math.copysign(Tc_mag, T_cl_dem)
            Jc = JI + (JE if k0 else 0.0)
            a_ref = R_tgt * a_w_prev
            T_gb_est = spin + (1 - eta) * abs(Tc)            # gearbox drag feed-forward
            Tm = Tc + T_gb_est - Te_c + Jc * (wi_tgt - wi) / max(p.shift_tau, 2 * dt) + Jc * a_ref
        Tm = min(max(Tm, -Tm_lim), Tm_lim)

        # battery power limits (virtual BMS) incl. HSG
        P_he = hsg.pelec(T_hsg, Gh * we) if (abs(T_hsg) > 0 or we > 1) and on_now else 0.0
        P_me = mot.pelec(Tm, wi)
        P_bus = P_me + P_he + P_aux
        P_bt = P_bus / eta_dc if P_bus > 0 else P_bus * eta_dc
        if P_bt > P_dis_lim and Tm > 0:
            Tm = mot.torque_for_pelec(P_dis_lim * eta_dc - P_he - P_aux, wi, -Tm_lim, Tm)
        elif P_bt < -P_chg_lim and Tm < 0:
            Tm = mot.torque_for_pelec(-P_chg_lim / eta_dc - P_he - P_aux, wi, Tm, Tm_lim)
        P_me = mot.pelec(Tm, wi)
        P_bus = P_me + P_he + P_aux
        P_bt = P_bus / eta_dc if P_bus > 0 else P_bus * eta_dc

        # gearbox losses (spin + load)
        T_through = abs(Tm + Te_c) if dct_lock else abs(Tc)
        T_gb = (spin + (1 - eta) * T_through) if (wi > 0.5 or T_through > 0) else 0.0
        if wi < 0.5:
            T_gb = min(T_gb, T_through) * 0.0

        # friction brakes take what the powertrain does not
        T_fb = 0.0
        if Tw_req < 0:
            T_pt_w = R_cl * Tc if not dct_lock else R_now * (Tm + Te_c - T_gb)
            T_fb = max(0.0, -Tw_req - max(0.0, -T_pt_w))
        elif brake > 0:
            T_fb = brake * T_brake_max

        # ================================================= PLANT
        mdot = eng.fuel_rate(we, Te) if fuel_on else 0.0
        I = bat.current(P_bt, soc)
        Voc, Rb = bat.ocv(soc), bat.r(soc)
        P_chem = Voc * I
        vrel = v + v_wind
        F_aero = cdA * vrel * abs(vrel)
        F_roll = p.Crr * p.mass * G * math.cos(th) if v > 0.01 else 0.0
        F_grade = p.mass * G * math.sin(th)
        T_fb_eff = T_fb if v > 0.01 else 0.0
        QE = Te_net + T_hsg * Gh
        QI = Tm - T_gb
        QV = -(F_aero + F_roll + F_grade) * rw - T_fb_eff
        R = R_now if dct_lock else R_cl
        if k0 and dct_lock:
            a_w = (R * (QE + QI) + QV) / ((JE + JI) * R * R + JV)
            a_i = a_e = R * a_w
        elif dct_lock:
            a_e = QE / JE
            a_w = (R * QI + QV) / (JI * R * R + JV)
            a_i = R * a_w
        elif k0:
            a_i = (QE + QI - Tc) / (JE + JI); a_e = a_i
            a_w = (QV + R * Tc) / JV
        else:
            a_e = QE / JE
            a_i = (QI - Tc) / JI
            a_w = (QV + R * Tc) / JV

        # ---- energy bookkeeping (start-of-step speeds)
        P_eng = Te_net * we
        P_m = Tm * wi
        P_hm = T_hsg * Gh * we
        E["fuel"] += mdot * LHV * dt
        E["eng_mech"] += P_eng * dt
        E["eng_loss"] += (mdot * LHV - P_eng) * dt
        E["mot_mech"] += P_m * dt; E["mot_elec"] += P_me * dt; E["mot_loss"] += (P_me - P_m) * dt
        if P_m >= 0:
            E["mot_motor_mech"] += P_m * dt; E["mot_motor_elec"] += P_me * dt
        else:
            E["mot_gen_mech"] += -P_m * dt; E["mot_gen_elec"] += -P_me * dt
        E["hsg_mech"] += P_hm * dt; E["hsg_elec"] += P_he * dt; E["hsg_loss"] += (P_he - P_hm) * dt
        E["dcdc_loss"] += (P_bt - P_bus) * dt
        E["batt_chem"] += P_chem * dt
        E["batt_loss"] += I * I * Rb * dt
        if P_chem >= 0:
            E["batt_out"] += P_chem * dt
        else:
            E["batt_in"] += -P_chem * dt
        E["aux"] += P_aux * dt
        E["aero"] += F_aero * v * dt
        E["roll"] += F_roll * v * dt
        E["grade"] += F_grade * v * dt
        E["brake"] += T_fb_eff * ww * dt
        E["gearbox"] += T_gb * wi * dt
        if not dct_lock:
            E["slip_dct"] += Tc * (wi - R * ww) * dt

        # mode classification
        if v < 0.3 and Tw_req <= 0 and eng_state != "CHARGE":
            mode = 0
        elif eng_state == "CHARGE":
            mode = 4
        elif k0:
            mode = 2 if Tm < -1.0 else 3
        else:
            mode = 1
        mode_time[mode] += dt
        gear_time[gear - 1] += dt
        if fuel_on and P_eng > 1000 and mdot > 0:
            eta_e = P_eng / (mdot * LHV)
            eng_eff.append(eta_e)
        else:
            eta_e = float("nan")

        if step % log_every == 0:
            vals = (t, v_r, v, accel, brake, mode, gear, 1.0 if shift_to else 0.0, 1.0 if k0 else 0.0,
                    ENG_STATES.index(eng_state), we, wi, Te_net, Tm, T_hsg, T_fb_eff, Tc,
                    P_eng, P_m, P_hm, P_me, P_he, P_bt,
                    (R * (Tm + Te_c - T_gb) if dct_lock else R * Tc) * ww,
                    soc, I, Voc - I * Rb, mdot, fuel_kg / p.fuel_rho, dist,
                    P_dis_lim, P_chg_lim, ecms_s, eta_e, grade)
            for name, val in zip(SIGNALS, vals):
                log[name].append(val)

        # ---- integrate with constraints
        KE_b = ke()
        Q_pow = QE * we + QI * wi + QV * ww - (Tc * (wi - R * ww) if not dct_lock else 0.0)
        we += a_e * dt
        wi += a_i * dt
        ww += a_w * dt
        a_w_prev = a_w
        clamp = False
        if ww < 0.0:
            ww = 0.0; clamp = True
            if dct_lock:
                wi = 0.0
                if k0:
                    we = 0.0
        if we < 0.0:
            we = 0.0; clamp = True
        if wi < 0.0:
            wi = 0.0; clamp = True
        # shift completion: on-coming clutch locks at synchronism
        if shift_to:
            slip_n = wi - R_g[shift_to - 1] * ww
            if abs(slip_n) < 2.0 or shift_t > 2.0 or (slip_n > 0) != (wi - R_g[shift_to - 1] * (ww - a_w * dt) - a_i * dt > 0):
                KE_pre = ke()
                gear, shift_to = shift_to, 0
                last_shift = t
                wi = R_g[gear - 1] * ww
                if k0:
                    we = wi
                E["lock"] += KE_pre - ke()
        if eng_state == "RUN" and k0 and abs(we - wi) > 1e-9:
            KE_pre = ke(); we = wi; E["lock"] += KE_pre - ke()
        KE_pred = KE_b + Q_pow * dt
        if clamp:
            E["standstill"] += max(0.0, KE_pred - ke())

        soc = min(max(soc - I * dt / bat.Q, 0.0), 1.0)
        fuel_kg += mdot * dt
        dist += v * dt
        if progress is not None and step % 2000 == 0:
            progress(step / n_steps)
        if scen.kind == "wot" and wot_done_t is not None and v < 0.1 and t > wot_done_t:
            break

    res = Results()
    res.p, res.scen = p, scen
    res.eng, res.mot, res.hsg, res.bat = eng, mot, hsg, bat
    res.log = {k: np.array(v) for k, v in log.items()}
    E["dKE"] = ke()
    res.E = E
    res.mode_time = mode_time
    res.gear_time = gear_time
    res.n_starts = n_starts
    res.n_shifts = n_shifts
    res.eng_eff_samples = np.array(eng_eff)
    res.t100, res.t130 = t100, wot_done_t
    res.R_g = R_g
    res.summary = summarize(res, fuel_kg, dist, soc)
    return res


def summarize(res, fuel_kg, dist, soc_end):
    p, E, lg = res.p, res.E, res.log
    km = dist / 1000.0
    fuel_L = fuel_kg / p.fuel_rho
    L100 = fuel_L / km * 100 if km > 0.05 else float("nan")
    eta_e = E["eng_mech"] / E["fuel"] if E["fuel"] > 0 else 0.35
    eta_path = max(eta_e, 0.2) * 0.88
    fuel_corr = fuel_kg + E["batt_chem"] / (eta_path * p.fuel_LHV_MJkg * 1e6)
    L100c = fuel_corr / p.fuel_rho / km * 100 if km > 0.05 else float("nan")
    src = E["fuel"] + E["batt_chem"]
    sinks = sum(E[k] for k in ["eng_loss", "mot_loss", "hsg_loss", "dcdc_loss", "batt_loss", "aux",
                                "aero", "roll", "grade", "brake", "gearbox", "slip_dct", "lock",
                                "standstill", "dKE"])
    thr = abs(E["fuel"]) + E["batt_out"] + E["batt_in"]
    err = lg["v"] - lg["v_ref"]
    mv = lg["v_ref"] > 0.5
    tot = sum(res.mode_time) or 1
    return dict(
        scenario=res.scen.name, duration_s=float(lg["t"][-1]) if len(lg["t"]) else 0.0,
        distance_km=km, fuel_L=fuel_L, fuel_L100km=L100, fuel_L100km_soc_corrected=L100c,
        mpg_US=235.215 / L100 if L100 and L100 > 0 else float("nan"),
        co2_g_km=fuel_kg * p.co2_per_kg * 1000 / km if km > 0.05 else float("nan"),
        soc_start=p.soc_init, soc_end=soc_end, engine_starts=res.n_starts, shifts=res.n_shifts,
        engine_on_pct=100 * (res.mode_time[2] + res.mode_time[3] + res.mode_time[4]) / tot,
        ev_pct=100 * res.mode_time[1] / tot,
        engine_avg_eff_pct=100 * eta_e,
        speed_err_rms_kmh=float(np.sqrt(np.mean(err[mv] ** 2)) * 3.6) if np.any(mv) else 0.0,
        speed_err_max_kmh=float(np.max(np.abs(err)) * 3.6) if len(err) else 0.0,
        energy_balance_err_pct=100 * (src - sinks) / max(thr, 1.0),
        regen_recovered_kWh=E["mot_gen_elec"] / 3.6e6,
        friction_brake_kWh=E["brake"] / 3.6e6,
        t_0_100_s=res.t100 if res.scen.kind == "wot" else None,
    )


def energy_table(res):
    E = res.E
    k = lambda j: j / 3.6e6
    return [
        ("Source", "Fuel chemical energy", k(E["fuel"])),
        ("Source", "Battery net discharge (chemical)", k(E["batt_chem"])),
        ("Loss", "Engine losses", k(E["eng_loss"])),
        ("Loss", "Motor losses", k(E["mot_loss"])),
        ("Loss", "HSG losses", k(E["hsg_loss"])),
        ("Loss", "Battery resistive losses", k(E["batt_loss"])),
        ("Loss", "DC-DC converter losses", k(E["dcdc_loss"])),
        ("Loss", "Gearbox (mesh + spin)", k(E["gearbox"])),
        ("Loss", "DCT clutch slip (shifts)", k(E["slip_dct"])),
        ("Loss", "Friction brakes", k(E["brake"])),
        ("Load", "Auxiliary load", k(E["aux"])),
        ("Load", "Aerodynamic drag", k(E["aero"])),
        ("Load", "Rolling resistance", k(E["roll"])),
        ("Load", "Grade (potential energy)", k(E["grade"])),
        ("Storage", "Kinetic energy change", k(E["dKE"])),
        ("Loss", "Clutch lock-up / standstill (numerical)", k(E["lock"] + E["standstill"])),
    ]


def efficiency_table(res):
    E = res.E
    r = lambda a, b: 100 * a / b if b > 1e-6 else float("nan")
    return [
        ("Engine plant (mech out / fuel)", r(E["eng_mech"], E["fuel"])),
        ("Motor motoring (mech / elec)", r(E["mot_motor_mech"], E["mot_motor_elec"])),
        ("Motor generating (elec / mech)", r(E["mot_gen_elec"], E["mot_gen_mech"])),
        ("Battery (1 - loss / throughput)", 100 - r(E["batt_loss"], E["batt_out"] + E["batt_in"])),
        ("Gearbox (1 - loss / road+brake energy)",
         100 - r(E["gearbox"], E["aero"] + E["roll"] + max(E["grade"], 0) + E["brake"] + E["gearbox"])),
        ("Road loads / fuel (tank-to-wheel)",
         r(E["aero"] + E["roll"] + max(E["grade"], 0), E["fuel"] + max(E["batt_chem"], 0))),
    ]


def export_results(res, outdir):
    os.makedirs(outdir, exist_ok=True)
    base = res.scen.name.split()[0].replace("/", "-")
    ts = os.path.join(outdir, f"kia_p2dct_{base}_timeseries.csv")
    with open(ts, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(SIGNALS)
        for row in zip(*[res.log[s] for s in SIGNALS]):
            w.writerow([f"{x:.6g}" for x in row])
    sm = os.path.join(outdir, f"kia_p2dct_{base}_energy_summary.csv")
    with open(sm, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Group", "Item", "Energy [kWh]"])
        for g, it, val in energy_table(res):
            w.writerow([g, it, f"{val:.5f}"])
        w.writerow([]); w.writerow(["Efficiency", "Value [%]"])
        for it, val in efficiency_table(res):
            w.writerow([it, f"{val:.2f}"])
        w.writerow([]); w.writerow(["Result", "Value"])
        for kk, vv in res.summary.items():
            w.writerow([kk, vv])
        w.writerow([]); w.writerow(["Gear", "Time [s]"])
        for g, tt in enumerate(res.gear_time, 1):
            w.writerow([g, f"{tt:.1f}"])
    out = [ts, sm]
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        s1 = wb.active; s1.title = "Energy summary"
        s1.append(["Group", "Item", "Energy [kWh]"])
        for g, it, val in energy_table(res):
            s1.append([g, it, round(val, 5)])
        s1.append([]); s1.append(["Efficiency", "", "Value [%]"])
        for it, val in efficiency_table(res):
            s1.append(["", it, None if math.isnan(val) else round(val, 2)])
        s1.append([]); s1.append(["Result", "", "Value"])
        for kk, vv in res.summary.items():
            s1.append(["", kk, vv if not isinstance(vv, float) else round(vv, 4)])
        s1.column_dimensions["B"].width = 40
        s2 = wb.create_sheet("Parameters")
        s2.append(["Parameter", "Value"])
        for kk, vv in asdict(res.p).items():
            s2.append([kk, vv])
        s3 = wb.create_sheet("Time series")
        s3.append(SIGNALS)
        for row in zip(*[res.log[s] for s in SIGNALS]):
            s3.append([None if (isinstance(x, float) and math.isnan(x)) else float(x) for x in row])
        xl = os.path.join(outdir, f"kia_p2dct_{base}_results.xlsx")
        wb.save(xl)
        out.append(xl)
    except ImportError:
        pass
    return out


def print_summary(res):
    s = res.summary
    print(f"\n=== {s['scenario']} (Kia P2 + 6DCT) ===")
    print(f"  distance            {s['distance_km']:.2f} km in {s['duration_s']:.0f} s")
    print(f"  fuel                {s['fuel_L']:.3f} L  ->  {s['fuel_L100km']:.2f} L/100 km")
    print(f"  SOC-corrected       {s['fuel_L100km_soc_corrected']:.2f} L/100 km   CO2 {s['co2_g_km']:.0f} g/km")
    print(f"  SOC                 {s['soc_start']:.3f} -> {s['soc_end']:.3f}")
    print(f"  engine on           {s['engine_on_pct']:.1f} % of time, {s['engine_starts']} starts,"
          f" avg eff {s['engine_avg_eff_pct']:.1f} %")
    print(f"  EV                  {s['ev_pct']:.1f} %   shifts {s['shifts']}")
    print(f"  speed tracking      RMS {s['speed_err_rms_kmh']:.2f} km/h, max {s['speed_err_max_kmh']:.2f} km/h")
    print(f"  regen recovered     {s['regen_recovered_kWh']:.3f} kWh, friction brakes {s['friction_brake_kWh']:.3f} kWh")
    print(f"  energy balance err  {s['energy_balance_err_pct']:.3f} %")
    print("  gear time [s]       " + "  ".join(f"{g}:{tt:.0f}" for g, tt in enumerate(res.gear_time, 1)))
    if s.get("t_0_100_s") is not None:
        print(f"  0-100 km/h          {s['t_0_100_s']:.1f} s")


def main(argv=None):
    ap = argparse.ArgumentParser(description="Kia-style P2 hybrid with 6-speed DCT")
    ap.add_argument("--cycle", default="WLTC")
    ap.add_argument("--csv", default="")
    ap.add_argument("--set", action="append")
    ap.add_argument("--export", default="")
    ap.add_argument("--list-params", action="store_true")
    a = ap.parse_args(argv)
    p = P2Params()
    names = {f.name for f in fields(P2Params)}
    for it in a.set or []:
        k, v = it.split("=", 1)
        if k not in names:
            raise SystemExit(f"unknown parameter '{k}'")
        setattr(p, k, float(v))
    if a.list_params:
        for f in fields(P2Params):
            print(f"  {f.name:<18s} = {getattr(p, f.name):<10g} {PARAM_HELP.get(f.name, '')}")
        return
    alias = {"WLTC": "WLTC class 3b", "UDDS": "UDDS (FTP-72)", "HWFET": "HWFET", "US06": "US06",
             "NEDC": "NEDC", "WOT": "WOT 0-130 km/h", "CONST": "Constant speed", "CSV": "Custom CSV"}
    cyc = alias.get(a.cycle.upper(), a.cycle)
    res = simulate(p, make_scenario(cyc, p, a.csv))
    print_summary(res)
    if a.export:
        for f in export_results(res, a.export):
            print("  wrote", f)
    return res


if __name__ == "__main__":
    main()
